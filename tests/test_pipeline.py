"""End-to-end Milestone 1 pipeline tests."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook

from profile_intelligence.bootstrap import build_container
from profile_intelligence.services.application import ApplicationService
from profile_intelligence.services.import_service import ImportService
from profile_intelligence.services.profile_service import ProfileService


def _write_sample_csv(path: Path) -> None:
    path.write_text(
        "id,name,email,organization,title,location\n"
        "1,Ada Lovelace,ada@example.com,Analytical Engines,Analyst,London\n"
        "2,Grace Hopper,grace@example.com,US Navy,Admiral,New York\n",
        encoding="utf-8",
    )


def _write_sample_xlsx(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(
        ["external_id", "name", "email", "company", "role", "city"]
    )
    sheet.append(
        ["x-1", "Alan Turing", "alan@example.com", "Bletchley", "Mathematician", "UK"]
    )
    workbook.save(path)


def test_csv_import_search_export_score(temp_root: Path) -> None:
    container = build_container(root_dir=temp_root)
    app = container.resolve(ApplicationService)
    app.start()

    csv_path = temp_root / "people.csv"
    _write_sample_csv(csv_path)

    summary = container.resolve(ImportService).import_path(csv_path)
    assert summary.created == 2
    assert summary.updated == 0

    # Re-import should update, not duplicate.
    summary2 = container.resolve(ImportService).import_path(csv_path)
    assert summary2.created == 0
    assert summary2.updated == 2

    profiles = container.resolve(ProfileService)
    assert profiles.count() == 2

    found = profiles.search("lovelace")
    assert len(found) == 1
    assert found[0].email == "ada@example.com"
    assert found[0].score is not None and found[0].score >= 25

    export_path = profiles.export_excel(temp_root / "exports" / "out.xlsx")
    assert export_path.exists()
    workbook = load_workbook(export_path)
    sheet = workbook.active
    assert sheet["A1"].value == "id"
    assert sheet.max_row == 3  # header + 2 profiles

    rescored = profiles.rescore_all()
    assert rescored == 2
    app.shutdown()


def test_excel_import(temp_root: Path) -> None:
    container = build_container(root_dir=temp_root)
    app = container.resolve(ApplicationService)
    app.start()

    xlsx_path = temp_root / "people.xlsx"
    _write_sample_xlsx(xlsx_path)

    summary = container.resolve(ImportService).import_path(xlsx_path)
    assert summary.plugin == "excel"
    assert summary.created == 1

    profiles = container.resolve(ProfileService).list_profiles()
    assert profiles[0].display_name == "Alan Turing"
    assert profiles[0].organization == "Bletchley"
    app.shutdown()


def test_builtin_importers_discovered(temp_root: Path) -> None:
    container = build_container(root_dir=temp_root)
    app = container.resolve(ApplicationService)
    app.start()
    names = {plugin.name for plugin in app.importers.list_plugins()}
    assert names == {"csv", "excel", "webarchive"}
    app.shutdown()
