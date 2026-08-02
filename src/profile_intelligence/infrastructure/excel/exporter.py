"""Export stored profiles to Excel workbooks."""

from __future__ import annotations

from collections.abc import Sequence
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from profile_intelligence.core.config import AppConfig
from profile_intelligence.core.exceptions import ExcelError
from profile_intelligence.core.logging import get_logger
from profile_intelligence.core.types import PathLike
from profile_intelligence.infrastructure.database.models import Profile

logger = get_logger(__name__)

_HEADERS: tuple[str, ...] = (
    "id",
    "external_id",
    "display_name",
    "email",
    "phone",
    "title",
    "organization",
    "location",
    "tags",
    "source",
    "notes",
    "score",
    "created_at",
    "updated_at",
)


class ExcelExporter:
    """Write profile rows to an ``.xlsx`` workbook."""

    def __init__(self, config: AppConfig) -> None:
        self._config = config

    def export(
        self,
        profiles: Sequence[Profile],
        output_path: PathLike | None = None,
        *,
        sheet_name: str | None = None,
    ) -> Path:
        """Export *profiles* to Excel and return the output path."""
        self._config.ensure_directories()
        destination = (
            Path(output_path)
            if output_path is not None
            else self._config.exports_dir / "profiles.xlsx"
        )
        if not destination.is_absolute():
            destination = (self._config.root_dir / destination).resolve()
        destination.parent.mkdir(parents=True, exist_ok=True)

        title = sheet_name or self._config.excel.default_sheet_name
        try:
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = title[:31] or "Profiles"

            header_font = Font(bold=True)
            for col_index, header in enumerate(_HEADERS, start=1):
                cell = worksheet.cell(row=1, column=col_index, value=header)
                cell.font = header_font

            for row_index, profile in enumerate(profiles, start=2):
                values = (
                    profile.id,
                    profile.external_id,
                    profile.display_name,
                    profile.email,
                    profile.phone,
                    profile.title,
                    profile.organization,
                    profile.location,
                    profile.tags,
                    profile.source,
                    profile.notes,
                    profile.score,
                    profile.created_at.isoformat() if profile.created_at else None,
                    profile.updated_at.isoformat() if profile.updated_at else None,
                )
                for col_index, value in enumerate(values, start=1):
                    worksheet.cell(row=row_index, column=col_index, value=value)

            for col_index in range(1, len(_HEADERS) + 1):
                letter = get_column_letter(col_index)
                worksheet.column_dimensions[letter].width = 18

            workbook.save(destination)
        except ExcelError:
            raise
        except Exception as exc:
            raise ExcelError(
                f"Failed to export profiles to {destination}",
                cause=exc,
            ) from exc

        logger.info("Exported %d profile(s) to %s", len(profiles), destination)
        return destination
